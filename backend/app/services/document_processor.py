from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path
from typing import Any

SUPPORTED_DOCUMENT_EXTENSIONS = {".txt", ".csv", ".pdf", ".docx", ".jpg", ".jpeg", ".png", ".tif", ".tiff"}
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".tif", ".tiff"}

COORD_RE = re.compile(r"(?<!\d)([-+]?(?:[1-8]?\d(?:\.\d+)?|90(?:\.0+)?))\s*,\s*([-+]?(?:1[0-7]\d(?:\.\d+)?|[1-9]?\d(?:\.\d+)?|180(?:\.0+)?))(?!\d)")
DATE_RE = re.compile(r"\b(?:\d{4}-\d{1,2}-\d{1,2}|\d{1,2}[/-]\d{1,2}[/-]\d{2,4}|\d{1,2}\s+(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\s+\d{2,4})\b", re.IGNORECASE)
ENTITY_RE = re.compile(r"\b([A-Z][a-z]+(?:\s+[A-Z][a-z]+){1,2})\b")
PLACE_HINT_RE = re.compile(r"\b(?:at|near|in|from|toward|towards|route|checkpoint|village|market|bridge|road)\s+([A-Z][A-Za-z' -]{2,40})", re.IGNORECASE)

KEYWORD_GROUPS = {
    "vehicles": ["vehicle", "truck", "pickup", "pick-up", "toyota", "motorbike", "motorcycle"],
    "checkpoints": ["checkpoint", "check point", "barrier", "roadblock"],
    "incidents": ["incident", "attack", "ambush", "theft", "kidnap", "shooting"],
    "groups": ["group", "militia", "armed men", "gang"],
    "weapons": ["weapon", "rifle", "ak", "gun", "grenade"],
    "locations": ["village", "market", "bridge", "road", "river", "route"],
}


def is_supported_document(path: Path) -> bool:
    return path.suffix.lower() in SUPPORTED_DOCUMENT_EXTENSIONS


def extract_document(path: Path) -> dict[str, Any]:
    suffix = path.suffix.lower()
    text = ""
    status = "unsupported"
    detail = ""

    try:
        if suffix == ".txt":
            text = _read_text(path)
            status = "extracted"
        elif suffix == ".csv":
            text = _read_csv(path)
            status = "extracted"
        elif suffix == ".pdf":
            text, detail = _read_pdf(path)
            status = "extracted" if text else "needs_library"
        elif suffix == ".docx":
            text, detail = _read_docx(path)
            status = "extracted" if text else "needs_library"
        elif suffix in IMAGE_EXTENSIONS:
            status = "ocr_not_configured"
            detail = "Image OCR is planned, but Tesseract is not configured in this MVP."
        else:
            detail = "File type is not supported by the current MVP extractor."
    except Exception as exc:
        status = "error"
        detail = f"{type(exc).__name__}: {str(exc)[:180]}"

    text = (text or "").strip()
    clues = extract_clues(text)
    if status == "extracted" and not text:
        status = "no_text_found"
    return {
        "status": status,
        "detail": detail,
        "text": text,
        "clues": clues,
    }


def extract_clues(text: str) -> dict[str, Any]:
    sample = text or ""
    lower = sample.lower()
    coords = []
    seen_coords = set()
    for match in COORD_RE.finditer(sample):
        lat = round(float(match.group(1)), 6)
        lon = round(float(match.group(2)), 6)
        key = (lat, lon)
        if key not in seen_coords:
            seen_coords.add(key)
            coords.append({"latitude": lat, "longitude": lon, "text": match.group(0)})

    dates = _unique(DATE_RE.findall(sample), limit=12)
    entities = [value for value in _unique(ENTITY_RE.findall(sample), limit=20) if value.lower() not in {"Central African", "Local Field"}]
    places = _unique([m.strip(" .,;:-") for m in PLACE_HINT_RE.findall(sample)], limit=16)
    keywords = []
    for group, terms in KEYWORD_GROUPS.items():
        hits = [term for term in terms if term in lower]
        if hits:
            keywords.append({"group": group, "matches": hits[:6]})

    return {
        "dates": dates,
        "coordinates": coords[:10],
        "possible_entities": entities[:12],
        "possible_places": places[:12],
        "keywords": keywords,
    }


def clues_to_json(clues: dict[str, Any]) -> str:
    return json.dumps(clues, ensure_ascii=False)


def _read_text(path: Path) -> str:
    data = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _read_csv(path: Path) -> str:
    raw = _read_text(path)
    reader = csv.reader(io.StringIO(raw))
    lines = []
    for idx, row in enumerate(reader):
        if idx >= 200:
            lines.append("[CSV truncated after 200 rows]")
            break
        lines.append(" | ".join(cell.strip() for cell in row))
    return "\n".join(lines)


def _read_pdf(path: Path) -> tuple[str, str]:
    try:
        from pypdf import PdfReader  # type: ignore
    except Exception:
        return "", "PDF extraction needs pypdf installed."
    reader = PdfReader(str(path))
    parts = []
    for page in reader.pages[:30]:
        parts.append(page.extract_text() or "")
    text = "\n".join(parts)
    # Tag Waypoint-origin PDFs so analysts can identify the source in the DB
    detail = "origine Waypoint" if "WAYPOINT CORE" in text[:600] else ""
    return text, detail


def _read_docx(path: Path) -> tuple[str, str]:
    try:
        from docx import Document  # type: ignore
    except Exception:
        return "", "DOCX extraction needs python-docx installed."
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables[:10]:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts), ""


def _unique(values: list[str], limit: int) -> list[str]:
    result = []
    seen = set()
    for value in values:
        cleaned = " ".join(str(value).split()).strip()
        key = cleaned.lower()
        if cleaned and key not in seen:
            seen.add(key)
            result.append(cleaned)
        if len(result) >= limit:
            break
    return result
